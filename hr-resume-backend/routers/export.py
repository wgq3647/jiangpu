"""
导出Excel相关接口
支持按不同归类方式导出匹配结果
"""
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
import io
import random
import string
from datetime import datetime

from database import get_db
from models import MatchResult, Resume, Position, Company, User
from auth_utils import get_current_user

router = APIRouter(prefix="/export", tags=["导出"])


def get_filtered_results(db: Session, company: str = "", position: str = "",
                         owner: str = "", passed: str = "", min_score: int = 0,
                         name: str = ""):
    """获取筛选后的匹配结果"""
    query = db.query(MatchResult, Resume, Position, Company).join(
        Resume, MatchResult.resume_id == Resume.id
    ).join(
        Position, MatchResult.position_id == Position.id
    ).join(
        Company, Position.company_id == Company.id
    )

    if company:
        query = query.filter(Company.name == company)
    if position:
        query = query.filter(Position.name == position)
    if min_score and min_score > 0:
        query = query.filter(MatchResult.total_score >= min_score)
    if passed == "true":
        query = query.filter(MatchResult.passed_hard == True)
    elif passed == "false":
        query = query.filter(MatchResult.passed_hard == False)
    if name:
        query = query.filter(Resume.candidate_name.like(f"%{name}%"))

    results = query.all()

    # 关联上传人信息
    resume_owner_map = {}
    resumes = db.query(Resume).all()
    for r in resumes:
        owner_name = "未分配"
        if r.owner_id:
            user = db.query(User).filter(User.id == r.owner_id).first()
            if user:
                owner_name = user.name or user.phone
        resume_owner_map[r.id] = owner_name

    # 整理数据
    data = []
    for mr, resume, position, company_obj in results:
        data.append({
            "id": mr.id,
            "candidate_name": resume.candidate_name or mr.resume_name or "未知",
            "position_name": position.name,
            "company_name": company_obj.name,
            "total_score": mr.total_score or 0,
            "passed_hard": mr.passed_hard or False,
            "match_method": mr.match_method or "rule",
            "interview_result": mr.interview_feedback.interview_result if mr.interview_feedback else "",
            "interview_scheduled": mr.interview_scheduled or False,
            "sent_to_company": mr.sent_to_company or False,
            "owner_name": resume_owner_map.get(resume.id, "未分配"),
            "upload_date": resume.created_at.strftime("%Y-%m-%d") if resume.created_at else "",
            "resume_id": resume.id,
        })

    # 按上传人筛选
    if owner:
        data = [d for d in data if d["owner_name"] == owner]

    # 按匹配分降序
    data.sort(key=lambda x: x["total_score"], reverse=True)
    return data


def style_header(cell):
    """设置表头样式"""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="CBD5E1")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_group_header(cell, level=1):
    """设置分组标题样式"""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    colors = {
        1: ("EEF2FF", "4F46E5"),
        2: ("F0FDF4", "10B981"),
        3: ("EFF6FF", "3B82F6"),
    }
    bg, fg = colors.get(level, ("F8FAFC", "64748B"))
    cell.font = Font(bold=True, color=fg, size=11)
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    thin = Side(border_style="thin", color="CBD5E1")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_data_cell(cell):
    """设置数据单元格样式"""
    from openpyxl.styles import Alignment, Border, Side
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="E2E8F0")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def export_list_view(ws, data):
    """列表视图导出"""
    from openpyxl.utils import get_column_letter

    headers = ["序号", "候选人", "目标岗位", "公司", "总分", "硬性条件", "匹配方式", "面试状态", "上传人", "上传日期", "是否已发送"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        style_header(cell)

    for row_idx, item in enumerate(data, 2):
        values = [
            row_idx - 1,
            item["candidate_name"],
            item["position_name"],
            item["company_name"],
            item["total_score"],
            "满足硬性条件" if item["passed_hard"] else "不满足硬性条件",
            "AI匹配" if item["match_method"] == "ai" else "规则降级",
            get_interview_status(item),
            item["owner_name"],
            item["upload_date"],
            "已发送" if item["sent_to_company"] else "未发送",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            style_data_cell(cell)

    # 设置列宽
    widths = [6, 12, 18, 18, 8, 14, 10, 12, 12, 12, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def get_interview_status(item):
    """获取面试状态文本"""
    if item["interview_result"] == "通过":
        return "面试通过"
    elif item["interview_result"] == "不通过":
        return "面试不通过"
    elif item["interview_scheduled"]:
        return "已安排面试"
    return "未安排"


def export_grouped_by_company(ws, data):
    """按公司分组导出：公司→岗位→招聘人→候选人"""
    from openpyxl.utils import get_column_letter

    headers = ["层级", "候选人", "目标岗位", "公司", "总分", "硬性条件", "匹配方式", "面试状态", "上传人", "上传日期"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        style_header(cell)

    current_row = 2

    # 按公司分组
    companies = {}
    for item in data:
        cname = item["company_name"]
        if cname not in companies:
            companies[cname] = {}
        pname = item["position_name"]
        if pname not in companies[cname]:
            companies[cname][pname] = {}
        oname = item["owner_name"]
        if oname not in companies[cname][pname]:
            companies[cname][pname][oname] = []
        companies[cname][pname][oname].append(item)

    for cname, positions in companies.items():
        # L1 公司标题行
        total_count = sum(len(owners) for pos in positions.values() for owners in pos.values())
        passed_count = sum(1 for pos in positions.values() for owners in pos.values() for item in owners if item["passed_hard"])
        cell = ws.cell(row=current_row, column=1, value=f"🏢 L1 公司：{cname}（共{total_count}条，{passed_count}满足硬性条件）")
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
        style_group_header(cell, 1)
        current_row += 1

        for pname, owners in positions.items():
            # L2 岗位标题行
            pos_count = sum(len(o) for o in owners.values())
            cell = ws.cell(row=current_row, column=1, value=f"    💼 L2 岗位：{pname}（共{pos_count}条）")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
            style_group_header(cell, 2)
            current_row += 1

            for oname, items in owners.items():
                # L3 招聘人标题行
                cell = ws.cell(row=current_row, column=1, value=f"        👤 L3 招聘人：{oname}（共{len(items)}条）")
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
                style_group_header(cell, 3)
                current_row += 1

                # 候选人数据行
                for idx, item in enumerate(items, 1):
                    values = [
                        f"            {idx}",
                        item["candidate_name"],
                        item["position_name"],
                        item["company_name"],
                        item["total_score"],
                        "满足硬性条件" if item["passed_hard"] else "不满足硬性条件",
                        "AI匹配" if item["match_method"] == "ai" else "规则降级",
                        get_interview_status(item),
                        item["owner_name"],
                        item["upload_date"],
                    ]
                    for col, value in enumerate(values, 1):
                        cell = ws.cell(row=current_row, column=col, value=value)
                        style_data_cell(cell)
                    current_row += 1

    # 设置列宽
    widths = [20, 12, 18, 18, 8, 14, 10, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_grouped_by_owner(ws, data):
    """按招聘人员分组导出：招聘人→公司→岗位→候选人"""
    from openpyxl.utils import get_column_letter

    headers = ["层级", "候选人", "目标岗位", "公司", "总分", "硬性条件", "匹配方式", "面试状态", "上传人", "上传日期"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        style_header(cell)

    current_row = 2

    # 按招聘人分组
    owners = {}
    for item in data:
        oname = item["owner_name"]
        if oname not in owners:
            owners[oname] = {}
        cname = item["company_name"]
        if cname not in owners[oname]:
            owners[oname][cname] = {}
        pname = item["position_name"]
        if pname not in owners[oname][cname]:
            owners[oname][cname][pname] = []
        owners[oname][cname][pname].append(item)

    for oname, companies in owners.items():
        # L1 招聘人标题行
        total_count = sum(len(positions) for comps in companies.values() for positions in comps.values())
        cell = ws.cell(row=current_row, column=1, value=f"👤 L1 招聘人：{oname}（共{total_count}条）")
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
        style_group_header(cell, 1)
        current_row += 1

        for cname, positions in companies.items():
            # L2 公司标题行
            comp_count = sum(len(p) for p in positions.values())
            cell = ws.cell(row=current_row, column=1, value=f"    🏢 L2 公司：{cname}（共{comp_count}条）")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
            style_group_header(cell, 2)
            current_row += 1

            for pname, items in positions.items():
                # L3 岗位标题行
                cell = ws.cell(row=current_row, column=1, value=f"        💼 L3 岗位：{pname}（共{len(items)}条）")
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
                style_group_header(cell, 3)
                current_row += 1

                # 候选人数据行
                for idx, item in enumerate(items, 1):
                    values = [
                        f"            {idx}",
                        item["candidate_name"],
                        item["position_name"],
                        item["company_name"],
                        item["total_score"],
                        "满足硬性条件" if item["passed_hard"] else "不满足硬性条件",
                        "AI匹配" if item["match_method"] == "ai" else "规则降级",
                        get_interview_status(item),
                        item["owner_name"],
                        item["upload_date"],
                    ]
                    for col, value in enumerate(values, 1):
                        cell = ws.cell(row=current_row, column=col, value=value)
                        style_data_cell(cell)
                    current_row += 1

    # 设置列宽
    widths = [20, 12, 18, 18, 8, 14, 10, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_grouped_by_date(ws, data):
    """按上传时间分组导出：日期→候选人"""
    from openpyxl.utils import get_column_letter

    headers = ["层级", "候选人", "目标岗位", "公司", "总分", "硬性条件", "匹配方式", "面试状态", "上传人", "上传日期"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        style_header(cell)

    current_row = 2

    # 按日期分组
    dates = {}
    for item in data:
        d = item["upload_date"] or "未知日期"
        if d not in dates:
            dates[d] = []
        dates[d].append(item)

    # 按日期降序
    for d in sorted(dates.keys(), reverse=True):
        items = dates[d]
        # L1 日期标题行
        passed_count = sum(1 for item in items if item["passed_hard"])
        cell = ws.cell(row=current_row, column=1, value=f"📅 L1 日期：{d}（共{len(items)}条，{passed_count}满足硬性条件）")
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
        style_group_header(cell, 1)
        current_row += 1

        # 候选人数据行
        for idx, item in enumerate(items, 1):
            values = [
                f"    {idx}",
                item["candidate_name"],
                item["position_name"],
                item["company_name"],
                item["total_score"],
                "满足硬性条件" if item["passed_hard"] else "不满足硬性条件",
                "AI匹配" if item["match_method"] == "ai" else "规则降级",
                get_interview_status(item),
                item["owner_name"],
                item["upload_date"],
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=current_row, column=col, value=value)
                style_data_cell(cell)
            current_row += 1

    # 设置列宽
    widths = [20, 12, 18, 18, 8, 14, 10, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


@router.get("/match-results")
def export_match_results(
    view_mode: str = Query("list", description="视图模式：list/grouped/groupedByOwner/groupedByDate"),
    company: str = Query("", description="公司筛选"),
    position: str = Query("", description="岗位筛选"),
    owner: str = Query("", description="上传人筛选"),
    passed: str = Query("", description="硬性条件筛选：true/false"),
    min_score: int = Query(0, description="最低分"),
    name: str = Query("", description="候选人姓名搜索"),
    db: Session = Depends(get_db),
    current_user: dict = Depends(get_current_user),
):
    """导出匹配结果Excel，按当前归类方式分组呈现"""
    from openpyxl import Workbook

    data = get_filtered_results(db, company, position, owner, passed, min_score, name)

    wb = Workbook()
    ws = wb.active

    view_names = {
        "list": "List View",
        "grouped": "By Company",
        "groupedByOwner": "By Recruiter",
        "groupedByDate": "By Date",
    }
    ws.title = view_names.get(view_mode, "Match Results")

    if view_mode == "list":
        export_list_view(ws, data)
    elif view_mode == "grouped":
        export_grouped_by_company(ws, data)
    elif view_mode == "groupedByOwner":
        export_grouped_by_owner(ws, data)
    elif view_mode == "groupedByDate":
        export_grouped_by_date(ws, data)
    else:
        export_list_view(ws, data)

    # 冻结首行
    ws.freeze_panes = "A2"

    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # 生成随机文件名（英语+数字组合，避免中文编码问题）
    random_str = ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))
    filename = f"match_results_{view_names.get(view_mode, 'list').lower().replace(' ', '_')}_{random_str}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename.encode('utf-8').decode('latin-1')}"}
    )
