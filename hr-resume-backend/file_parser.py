"""
文件解析器
支持PDF、Word(docx)、图片、文本等多种简历格式的文本提取
"""
import os
import logging
from typing import Tuple, Optional

logger = logging.getLogger(__name__)

# PDF和图片识别功能开关配置键
PDF_IMAGE_RECOGNITION_CONFIG_KEY = "pdf_image_recognition_enabled"


def is_pdf_image_recognition_enabled(db) -> bool:
    """
    检查PDF和图片识别功能是否启用
    从数据库system_configs表读取配置，默认关闭
    """
    try:
        from models import SystemConfig
        config = db.query(SystemConfig).filter(
            SystemConfig.config_key == PDF_IMAGE_RECOGNITION_CONFIG_KEY
        ).first()
        if config and config.config_value:
            return config.config_value.lower() in ("true", "1", "yes", "on")
    except Exception as e:
        logger.warning(f"读取PDF/图片识别配置失败: {e}，默认关闭")
    return False


class FileParseError(Exception):
    """文件解析异常"""
    pass


class FileParser:
    """
    多模态文件解析器
    支持：PDF、Word(docx)、图片(png/jpg)、文本(txt/md)
    """

    # 支持的文件扩展名
    SUPPORTED_EXTENSIONS = {
        '.pdf': 'pdf',
        '.docx': 'word',
        '.doc': 'word',
        '.txt': 'text',
        '.md': 'text',
        '.png': 'image',
        '.jpg': 'image',
        '.jpeg': 'image',
        '.gif': 'image',
        '.bmp': 'image',
    }

    @staticmethod
    def get_file_type(file_name: str) -> Optional[str]:
        """根据文件名获取文件类型"""
        ext = os.path.splitext(file_name)[1].lower()
        return FileParser.SUPPORTED_EXTENSIONS.get(ext)

    @staticmethod
    def is_supported(file_name: str) -> bool:
        """检查文件是否支持解析"""
        return FileParser.get_file_type(file_name) is not None

    @staticmethod
    def parse(file_bytes: bytes, file_name: str) -> Tuple[str, str]:
        """
        解析文件，提取文本内容
        Args:
            file_bytes: 文件二进制内容
            file_name: 文件名（用于判断类型）
        Returns:
            (提取的文本, 文件类型)
        Raises:
            FileParseError: 解析失败时抛出
        """
        file_type = FileParser.get_file_type(file_name)
        if not file_type:
            raise FileParseError(f"不支持的文件格式: {file_name}")

        if file_type == 'pdf':
            return FileParser._parse_pdf(file_bytes), 'pdf'
        elif file_type == 'word':
            return FileParser._parse_word(file_bytes, file_name), 'word'
        elif file_type == 'image':
            return FileParser._parse_image(file_bytes, file_name), 'image'
        elif file_type == 'text':
            return FileParser._parse_text(file_bytes), 'text'
        else:
            raise FileParseError(f"未知文件类型: {file_type}")

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        """解析PDF文件，提取文本"""
        try:
            import io
            from PyPDF2 import PdfReader

            pdf_file = io.BytesIO(file_bytes)
            reader = PdfReader(pdf_file)

            texts = []
            for page_num, page in enumerate(reader.pages):
                try:
                    text = page.extract_text()
                    if text and text.strip():
                        texts.append(text.strip())
                except Exception as e:
                    logger.warning(f"PDF第{page_num+1}页提取失败: {e}")
                    continue

            if not texts:
                raise FileParseError("PDF中未提取到文本（可能是扫描件，需要OCR）")

            full_text = '\n'.join(texts)
            logger.info(f"PDF解析成功，共{len(reader.pages)}页，{len(full_text)}字符")
            return full_text

        except ImportError:
            raise FileParseError("PDF解析库未安装，请安装PyPDF2")
        except FileParseError:
            raise
        except Exception as e:
            raise FileParseError(f"PDF解析失败: {str(e)}")

    @staticmethod
    def _parse_word(file_bytes: bytes, file_name: str) -> str:
        """解析Word文件(docx)，提取文本"""
        ext = os.path.splitext(file_name)[1].lower()

        if ext == '.docx':
            try:
                import io
                from docx import Document

                doc_file = io.BytesIO(file_bytes)
                doc = Document(doc_file)

                texts = []
                # 提取段落
                for para in doc.paragraphs:
                    if para.text and para.text.strip():
                        texts.append(para.text.strip())

                # 提取表格内容
                for table in doc.tables:
                    for row in table.rows:
                        row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_texts:
                            texts.append(' | '.join(row_texts))

                if not texts:
                    raise FileParseError("Word文档中未提取到文本")

                full_text = '\n'.join(texts)
                logger.info(f"Word解析成功，{len(full_text)}字符")
                return full_text

            except ImportError:
                raise FileParseError("Word解析库未安装，请安装python-docx")
            except Exception as e:
                raise FileParseError(f"Word解析失败: {str(e)}")
        else:
            # .doc 旧格式，python-docx不支持，提示
            raise FileParseError("旧版.doc格式暂不支持，请另存为.docx格式后上传")

    @staticmethod
    def _parse_image(file_bytes: bytes, file_name: str) -> str:
        """
        解析图片简历
        优先使用OCR，如未安装OCR则返回提示信息
        （生产环境可接入AI视觉模型或Tesseract OCR）
        """
        # 尝试使用pytesseract OCR
        try:
            import io
            from PIL import Image
            import pytesseract

            img = Image.open(io.BytesIO(file_bytes))
            # 尝试中文+英文识别
            text = pytesseract.image_to_string(img, lang='chi_sim+eng')
            if text and text.strip():
                logger.info(f"图片OCR解析成功，{len(text)}字符")
                return text.strip()
            else:
                raise FileParseError("图片OCR未识别到文字")
        except ImportError:
            # OCR未安装，返回占位提示
            logger.warning("OCR库未安装，图片简历将使用占位文本")
            return f"[图片简历:{file_name}]\n（图片OCR功能未启用，如需解析图片简历请安装pytesseract和Tesseract引擎）"
        except Exception as e:
            logger.warning(f"图片OCR失败: {e}，使用占位文本")
            return f"[图片简历:{file_name}]\n（图片解析失败: {str(e)}）"

    @staticmethod
    def _parse_text(file_bytes: bytes) -> str:
        """解析纯文本文件"""
        # 尝试多种编码
        for encoding in ['utf-8', 'gbk', 'gb2312', 'utf-16']:
            try:
                text = file_bytes.decode(encoding)
                if text.strip():
                    return text.strip()
            except (UnicodeDecodeError, UnicodeError):
                continue
        # 全部失败，用忽略错误的方式
        return file_bytes.decode('utf-8', errors='ignore').strip()
